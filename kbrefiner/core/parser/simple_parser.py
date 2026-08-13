"""轻量文档解析器（不依赖 MinerU）。

基于 pypdf + python-docx 实现基础文档解析：
- PDF: pypdf 提取文本，按页分割
- DOCX: python-docx 提取段落和表格

适用场景：
- MinerU 未安装时的回退方案
- 文档结构简单（纯文本为主，无复杂表格/公式）
- 快速预览和测试

限制：
- 不支持 OCR（ scanned PDF / 图片需 MinerU）
- 不支持公式转 LaTeX
- 表格提取能力有限（仅 DOCX 原生表格）
- 版面还原能力弱
"""
from __future__ import annotations

import logging
from pathlib import Path

from kbrefiner.core.parser.base import (
    BaseParser,
    EmptyDocumentError,
    FileType,
    Heading,
    Table,
    Formula,
    ParseFailedError,
    ParsedDocument,
    UnsupportedFileError,
)

logger = logging.getLogger(__name__)


class SimpleParser(BaseParser):
    """轻量文档解析器，基于 pypdf + python-docx。"""

    @property
    def supported_types(self) -> set[FileType]:
        return {FileType.PDF, FileType.DOCX}

    def parse(self, file_path: str | Path) -> ParsedDocument:
        """解析文档，返回 ParsedDocument。"""
        path = Path(file_path)
        file_type = FileType.from_path(path)

        if not self.supports(file_type):
            raise UnsupportedFileError(f"SimpleParser 不支持此格式: {file_type}")

        if not path.exists():
            raise ParseFailedError(f"文件不存在: {path}")

        if file_type == FileType.PDF:
            markdown = self._parse_pdf(path)
        elif file_type == FileType.DOCX:
            markdown = self._parse_docx(path)
        else:
            raise UnsupportedFileError(f"不支持的格式: {file_type}")

        if not markdown.strip():
            raise EmptyDocumentError(f"解析结果为空文档: {path}")

        headings = self._extract_headings(markdown)
        tables = self._extract_tables(markdown)

        return ParsedDocument(
            source=str(path),
            file_type=file_type,
            markdown=markdown,
            headings=headings,
            tables=tables,
            page_count=self._detect_page_count(markdown),
        )

    def _parse_pdf(self, path: Path) -> str:
        """使用 pypdf 解析 PDF。"""
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ParseFailedError(
                "pypdf 未安装。请运行: pip install pypdf"
            )

        try:
            reader = PdfReader(str(path))
            parts: list[str] = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    parts.append(f"<!-- page {i + 1} -->\n{text}")
            return "\n\n".join(parts)
        except Exception as e:
            raise ParseFailedError(f"PDF 解析失败: {e}") from e

    def _parse_docx(self, path: Path) -> str:
        """使用 python-docx 解析 DOCX。"""
        try:
            import docx
        except ImportError:
            raise ParseFailedError(
                "python-docx 未安装。请运行: pip install python-docx"
            )

        try:
            doc = docx.Document(str(path))
            parts: list[str] = []

            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                # 根据样式映射标题层级
                style_name = (para.style.name or "").lower()
                if style_name.startswith("heading"):
                    try:
                        level = int(style_name.replace("heading", "").strip())
                    except ValueError:
                        level = 1
                    parts.append(f"{'#' * level} {text}")
                else:
                    parts.append(text)

            # 提取表格为 Markdown 格式
            for table in doc.tables:
                parts.append(self._table_to_markdown(table))

            return "\n\n".join(parts)
        except Exception as e:
            raise ParseFailedError(f"DOCX 解析失败: {e}") from e

    @staticmethod
    def _table_to_markdown(table) -> str:
        """将 python-docx 表格转为 Markdown 表格。"""
        import re

        def clean_cell(text: str) -> str:
            return re.sub(r"\|", "\\|", text.strip()).replace("\n", " ")

        rows = []
        for row in table.rows:
            cells = [clean_cell(cell.text) for cell in row.cells]
            rows.append(f"| {' | '.join(cells)} |")

        if not rows:
            return ""

        # 在第一行后添加分隔行
        header = rows[0]
        col_count = header.count("|") - 1
        separator = f"| {' | '.join(['---'] * col_count)} |"
        return f"{header}\n{separator}\n" + "\n".join(rows[1:])

    @staticmethod
    def _extract_headings(markdown: str) -> list[Heading]:
        """从 Markdown 提取标题层级。"""
        import re

        heading_re = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
        headings: list[Heading] = []
        for i, line in enumerate(markdown.splitlines()):
            m = heading_re.match(line)
            if m:
                headings.append(Heading(
                    level=len(m.group(1)),
                    text=m.group(2).strip(),
                    line_start=i,
                    line_end=i,
                ))
        return headings

    @staticmethod
    def _extract_tables(markdown: str) -> list[Table]:
        """提取 Markdown 表格（简单实现）。"""
        lines = markdown.splitlines()
        tables: list[Table] = []
        i = 0
        while i < len(lines):
            if "|" in lines[i] and i + 1 < len(lines) and "---" in lines[i + 1]:
                # 找到表格起始
                start = i
                i += 2  # 跳过分隔行
                while i < len(lines) and "|" in lines[i]:
                    i += 1
                table_text = "\n".join(lines[start:i])
                tables.append(Table(
                    content_html=table_text,
                    line_start=start,
                    line_end=i - 1,
                ))
            else:
                i += 1
        return tables

    @staticmethod
    def _detect_page_count(markdown: str) -> int:
        """从 page 标记检测页数。"""
        import re
        markers = re.findall(r"<!--\s*page\s*(\d+)\s*-->", markdown, re.IGNORECASE)
        return len(set(markers)) if markers else 0


__all__ = ["SimpleParser"]
